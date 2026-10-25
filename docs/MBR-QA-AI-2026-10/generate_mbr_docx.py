"""Generate MBR QA + AI Word document."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from datetime import date

OUT = r"d:\D-Drive\Personal\My-agent\docs\MBR-QA-AI-2026-10\MBR-QA-and-AI-Engagement-Platform.docx"


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()


def main():
    doc = Document()
    title = doc.add_heading("MBR: QA in AI — Engagement Platform & Delivery", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.add_run(f"Prepared: {date.today().strftime('%d %B %Y')}  |  Audience: Pravin Srivastava MBR  |  Contributors: QA + Engineering (Abhinav Sharma perspective)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "AI is delivering real gains on PR review, release reporting, and test-artefact hygiene. "
        "It is not replacing deep product and business validation. Delivery volume has increased; "
        "validation effort and headcount pressure have not fallen proportionally. Domain knowledge, "
        "Zephyr coverage, and distributed MCP/rules governance are the main constraints."
    )

    doc.add_heading("Slide 1 — What is running well", level=1)
    add_table(
        doc,
        ["Theme", "What is working", "Example / reference"],
        [
            ("Automated code review", "PR checks before merge", "Automation script review on PRs"),
            ("Test management + evidence", "Automation updates status/evidence in test tool", "Zephyr + automation pipelines"),
            ("AI-assisted release reporting", "Faster release notes / status", "Release reports generated using AI"),
            ("AI-assisted test hygiene", "Labels, priority, evidence, bug links reviewed", "Functional/regression review by AI"),
            ("Delivery throughput", "High story volume", "EP stories created — see Reference data"),
            ("Traceability (partial)", "Test Case Linked on Jira stories", "Jan–Jun 2026 — see Reference data"),
        ],
    )
    doc.add_paragraph("Projects to highlight: Engagement Platform (EP), Novo/PAP/re-enrollment, NeoEngage migrations, EOD4 retesting (e.g. cycle EP-R232).")

    doc.add_heading("Slide 2 — Where we are lagging", level=1)
    add_table(
        doc,
        ["Gap", "Symptom", "Evidence"],
        [
            ("Product/business knowledge", "Plausible but wrong AI output", "Major challenge — validation of new information"),
            ("Validation tax", "High script/TC volume from AI", "QA sub-tasks vs stories — Reference data"),
            ("Process uniformity", "Not all squads on same AI+MCP stack", "Multiple AI tools in same project"),
            ("MCP/rules silos", "Updates only by core people", "Others reluctant, unaware, or not passionate"),
            ("Lateral hires", "Misaligned with AI-first workflow", "Onboarding/playbooks gap"),
            ("Central monitoring", "Micro-issues not captured centrally", "Central team monitoring — not fully established"),
            ("Headcount model", "Volume up; FTE not down", "Human labour count cannot reduce as work increased"),
            ("AI used as-is", "Defects escape", "Hypercare/production bugs (e.g. attachment/multi-product scenarios)"),
            ("Zephyr coverage holes", "Cycle with zero cases", "RE-enrollment Retesting EOD4 (EP-R232): Test cases 0"),
            ("Metrics gap", "Cannot reproduce HR count vs AI work", "Zephyr export + tooling (e.g. ADO CLI) needed for MBR data"),
        ],
    )

    doc.add_heading("Slide 3 — AI challenges (consolidated)", level=1)
    challenges = [
        "Domain depth — product and business knowledge is the bottleneck; AI multiplies text, not truth.",
        "Validation tax — generated scripts/TCs/docs increase; creation savings < verification cost.",
        "Process uniformity — not all projects on one AI + MCP + rules stack.",
        "MCP adoption reluctance — cultural and skill, not only tooling.",
        "Knowledge silos — rules/MCP depend on few champions.",
        "Hiring fit — lateral hires not trained for evidence-labelled AI-assisted delivery.",
        "Governance — need central QA/AI observability (micro-issues, misuse, passed-as-is incidents).",
        "Headcount — volume increased; automation shifted work type, did not remove scrutiny.",
        "Quality risk — unreviewed AI in DB, workflows, test packs → production bugs.",
        "Leadership — people managers must become AI-augmented leaders with solid business depth.",
        "Tool sprawl — multiple AI tools per project without one standard.",
        "Fear/avoidance — some tech leads/engineers avoid AI.",
        "Enterprise license — ROI unclear without usage and quality metrics.",
        "Distribution — central AI not embedded in every project onboarding.",
        "Self-learning (Abhinav) — rules/skills should grow from team work, not only central suggestion.",
    ]
    for i, c in enumerate(challenges, 1):
        doc.add_paragraph(f"{i}. {c}", style="List Number")

    doc.add_heading("Abhinav Sharma — detailed view", level=1)
    doc.add_heading("Current state (honest)", level=2)
    doc.add_paragraph(
        "AI value is real but uneven: strong on PR review, release text, test artefact hygiene; "
        "weak where business rules and multi-step flows matter (PAP, multi-product, re-enrollment, attachments). "
        "Rules/MCP/skills are a product — if only a few maintain them, the org scales people, not capability."
    )
    doc.add_heading("Desired: self-learning mode (with guardrails)", level=2)
    doc.add_paragraph(
        "Self-learning is not raw chat dump. It requires: (1) shared repo for rules/skills, "
        "(2) mandatory write-back after validated work, (3) lightweight review (like PR) for rule changes, "
        "(4) sanitized patterns only — privacy/compliance. Anyone working in AI should update skills/rules "
        "from their work; leading teams keep adding without waiting for central suggestion only."
    )
    add_table(
        doc,
        ["Layer", "Owner", "What good looks like"],
        [
            ("Central", "Platform / QA CoE", "MCP catalog, enterprise license decision, metrics, micro-issue register"),
            ("Project", "Tech lead + QA lead", "One AI stack, Zephyr linking, onboarding checklist"),
            ("Individual", "Engineer / QA", "Evidence labels; skill/rule PR when pattern repeats ≥3 times"),
            ("Self-learning", "Automated + peer", "Weekly: top corrected AI mistakes → draft skill diff → TL approves"),
        ],
    )

    doc.add_heading("Reference data (Jira EP — 19 Oct 2026)", level=1)
    doc.add_paragraph("Periods as agreed:")
    add_table(
        doc,
        ["Period", "Date range"],
        [
            ("Period 1", "1 Jan 2026 – 31 Mar 2026"),
            ("Period 2", "1 Apr 2026 – 30 Jun 2026"),
            ("Period 3", "1 Jul 2026 – 15 Oct 2026"),
        ],
    )
    add_table(
        doc,
        ["Metric", "P1", "P2", "P3", "Source"],
        [
            ("EP Stories created", "583", "784", "1,280", "Jira EP"),
            ("QA test-case / UI validation sub-tasks", "9", "171", "223", "Jira EP"),
            ('Stories "Test Case Linked = Yes" (created in period)', "88", "124", "4", "Jira EP — P3 undercounts if link set later"),
            ("Avg TC/story by Functional/Regression/Sanity/Performance/Security", "N/A", "N/A", "N/A", "Zephyr export required"),
            ("Re-enrollment retest cycle EP-R232", "—", "—", "0 test cases in cycle", "Zephyr (screenshot)"),
        ],
    )
    doc.add_paragraph(
        "Interpretation: delivery accelerated; QA creation activity rose; executable Zephyr coverage can still be zero "
        "on critical cycles — AI and automation did not remove human validation; risk moves to 'did we run the right cases?'"
    )

    doc.add_heading("Next-line action items", level=1)
    doc.add_heading("A. QA perspective", level=2)
    add_table(
        doc,
        ["#", "Action", "Owner"],
        [
            ("1", "Mandatory Zephyr link before QA Successful for P1/P2 bugs", "QA lead EP"),
            ("2", "Fill EP-R232 + Novo PAP multi-product TCs (incl. EP-19286 repro)", "QA + BA"),
            ("3", "Monthly Zephyr export: TC count by type / linked story → MBR", "QA CoE"),
            ("4", "AI test review checklist on sample 10% (label, priority, evidence, bug link)", "QA"),
            ("5", 'Log "AI passed as-is" incidents in central register', "QA + TL"),
        ],
    )
    doc.add_heading("B. Engineering / AI governance", level=2)
    add_table(
        doc,
        ["#", "Action", "Owner"],
        [
            ("1", "One AI tool standard per project (+ exception process)", "Eng director"),
            ("2", "Project onboarding: MCP, rules, skills repo, evidence standard", "Central AI + PMO"),
            ("3", "Rule/skill PR workflow — any team member proposes; TL merges", "Platform"),
            ("4", "Enterprise license pilot with usage + defect escape metrics", "Finance + Eng"),
            ("5", "AI-delivered leader profile: business depth + AI workflow", "HR + leadership"),
            ("6", "Central micro-issue dashboard", "QA CoE"),
        ],
    )
    doc.add_heading("Enterprise license — decision frame", level=2)
    doc.add_paragraph(
        "Decide on enterprise when tracking: (1) % engineers active weekly, (2) time saved on release/report/PR, "
        "(3) validation hours trend, (4) prod escape / hypercare P1 rate. Standardize process first; "
        "license scales access, not quality."
    )

    doc.add_heading("Closing line for MBR", level=1)
    doc.add_paragraph(
        "AI saves draft and review time on code and documents, but it does not replace business-validated test depth. "
        "When product knowledge is not in rules, Zephyr, and people, output volume goes up while escape risk stays the same."
    )

    doc.add_paragraph()
    p = doc.add_paragraph("Document generated from MBR discussion notes (Pravin Srivastava / Abhinav Sharma).")
    p.runs[0].font.size = Pt(9)
    p.runs[0].italic = True

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
